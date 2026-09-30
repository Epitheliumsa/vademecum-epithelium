# Vista previa aproximada de las diapositivas (HTML + Chromium), porque LibreOffice no funciona en este entorno.
# Dibuja fondos, imágenes, formas, textos y tablas en su posición real, con Carlito (métricas de Calibri).
import sys, base64, os, html, subprocess
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

S = os.path.dirname(os.path.abspath(__file__))
src = sys.argv[1]; pref = sys.argv[2] if len(sys.argv) > 2 else 'prev'
solo = set(int(x) for x in sys.argv[3].split(',')) if len(sys.argv) > 3 else None
prs = Presentation(src)
EMU_PX = 96 / 914400
px = lambda e: e * EMU_PX
W, H = px(prs.slide_width), px(prs.slide_height)

def data_uri(blob, ct): return f'data:{ct};base64,' + base64.b64encode(blob).decode()

def fondo(slide):
    for part_el, part in ((slide._element, slide.part), (slide.slide_layout._element, slide.slide_layout.part), (slide.slide_layout.slide_master._element, slide.slide_layout.slide_master.part)):
        bg = part_el.find('.//' + qn('p:bg'))
        if bg is not None:
            blip = bg.find('.//' + qn('a:blip'))
            if blip is not None:
                p = part.related_part(blip.get(qn('r:embed')))
                return f'background:url({data_uri(p.blob, p.content_type)}) center/100% 100%;'
    return 'background:#fff;'

def color_de(el):
    if el is None: return None
    s = el.find(qn('a:srgbClr'))
    if s is not None:
        a = s.find(qn('a:alpha'))
        op = int(a.get('val')) / 100000 if a is not None else 1
        h = s.get('val'); return f'rgba({int(h[0:2],16)},{int(h[2:4],16)},{int(h[4:6],16)},{op})'
    if el.find(qn('a:schemeClr')) is not None:
        v = el.find(qn('a:schemeClr')).get('val')
        return {'bg1': '#fff', 'lt1': '#fff', 'tx1': '#000', 'dk1': '#000', 'accent1': '#4472C4'}.get(v, '#888')
    return None

def texto_html(tf, def_size=18, def_color='#000', lstyle=None):
    out = ''
    for p in tf.paragraphs:
        al = {1: 'left', 2: 'center', 3: 'right'}.get(p.alignment if p.alignment is None else int(p.alignment), 'left') if p.alignment else 'left'
        pPr = p._p.find(qn('a:pPr'))
        bullet = ''
        if pPr is not None and pPr.find(qn('a:buAutoNum')) is not None: bullet = '#. '
        runs = ''
        for r in p._p:
            if r.tag == qn('a:br'): runs += '<br>'; continue
            if r.tag != qn('a:r'): continue
            rPr = r.find(qn('a:rPr'))
            sz = int(rPr.get('sz')) / 100 if rPr is not None and rPr.get('sz') else def_size
            b = rPr is not None and rPr.get('b') == '1'
            col = color_de(rPr.find(qn('a:solidFill'))) if rPr is not None and rPr.find(qn('a:solidFill')) is not None else def_color
            t = r.find(qn('a:t')).text or ''
            runs += f'<span style="font-size:{sz*96/72:.1f}px;font-weight:{700 if b else 400};color:{col}">{html.escape(t)}</span>'
        sa = p.space_after.pt if p.space_after is not None else 0
        out += f'<div style="text-align:{al};margin-bottom:{sa*96/72:.1f}px;line-height:1.12">{bullet}{runs or "&nbsp;"}</div>'
    return out

def caja_texto(sh, def_size=18, def_color='#000', extra=''):
    tf = sh.text_frame
    bp = tf._txBody.find(qn('a:bodyPr'))
    ml = px(tf.margin_left); mr = px(tf.margin_right); mt = px(tf.margin_top); mb = px(tf.margin_bottom)
    anchor = {'t': 'flex-start', 'ctr': 'center', 'b': 'flex-end'}.get(bp.get('anchor') if bp is not None else 't', 'flex-start')
    return (f'<div style="position:absolute;left:{px(sh.left)}px;top:{px(sh.top)}px;width:{px(sh.width)}px;height:{px(sh.height)}px;'
            f'padding:{mt}px {mr}px {mb}px {ml}px;box-sizing:border-box;display:flex;flex-direction:column;justify-content:{anchor};{extra}">'
            f'<div>{texto_html(tf, def_size, def_color)}</div></div>')

def shape_html(sh, slide):
    if sh.left is None: return ''
    x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
    t = sh.shape_type
    if t == MSO_SHAPE_TYPE.PICTURE or t == MSO_SHAPE_TYPE.MEDIA:
        blip = sh._element.find('.//' + qn('a:blip'))
        rid = blip.get(qn('r:embed')) if blip is not None else None
        if not rid:
            svg = sh._element.find('.//{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip')
            rid = svg.get(qn('r:embed')) if svg is not None else None
        if not rid: return ''
        p = slide.part.related_part(rid)
        return f'<img src="{data_uri(p.blob, p.content_type)}" style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px">'
    if t == MSO_SHAPE_TYPE.TABLE:
        tb = sh.table; filas = ''
        for r in tb.rows:
            celdas = ''
            for j, c in enumerate(r.cells):
                tcPr = c._tc.find(qn('a:tcPr'))
                bg = color_de(tcPr.find(qn('a:solidFill'))) if tcPr is not None and tcPr.find(qn('a:solidFill')) is not None else '#fff'
                celdas += (f'<td style="width:{px(tb.columns[j].width)}px;background:{bg};padding:{px(c.margin_top)}px {px(c.margin_right)}px;'
                           f'vertical-align:middle;border:1px solid #fff;box-sizing:border-box">{texto_html(c.text_frame, 18)}</td>')
            filas += f'<tr style="height:{px(r.height)}px">{celdas}</tr>'
        return f'<table style="position:absolute;left:{x}px;top:{y}px;width:{w}px;border-collapse:collapse;table-layout:fixed">{filas}</table>'
    if t == MSO_SHAPE_TYPE.GROUP: return ''
    # formas y textos
    spPr = sh._element.find(qn('p:spPr'))
    estilo = ''
    if spPr is not None:
        f = spPr.find(qn('a:solidFill'))
        if f is not None: estilo += f'background:{color_de(f)};'
        ln = spPr.find(qn('a:ln'))
        if ln is not None and ln.find(qn('a:solidFill')) is not None:
            wln = int(ln.get('w', 12700)) * EMU_PX
            estilo += f'outline:{wln:.1f}px solid {color_de(ln.find(qn("a:solidFill")))};outline-offset:-{wln/2:.1f}px;'
        geom = spPr.find(qn('a:prstGeom'))
        g = geom.get('prst') if geom is not None else 'rect'
        if g == 'ellipse': estilo += 'border-radius:50%;'
        if g == 'roundRect': estilo += f'border-radius:{min(w,h)*0.12}px;'
    if sh.is_placeholder:
        pt = sh.placeholder_format.type
        size = 44 if 'TITLE' in str(pt) else 20
        if str(pt).startswith('CENTER_TITLE'): size = 54
        if str(pt).startswith('SUBTITLE'): size = 24
        return caja_texto(sh, size, '#000', estilo)
    if sh.has_text_frame and sh.text_frame.text.strip():
        return caja_texto(sh, 18, '#000', estilo)
    return f'<div style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;{estilo}"></div>'

paginas = []
for i, s in enumerate(prs.slides, 1):
    if solo and i not in solo: continue
    cuerpo = ''.join(shape_html(sh, s) for sh in s.shapes)
    paginas.append((i, f'<div class="sl" style="{fondo(s)}">{cuerpo}<div class="n">{i}</div></div>'))
doc = ('<!doctype html><html><head><meta charset="utf-8"><style>body{margin:0;font-family:Carlito,sans-serif}'
       f'.sl{{position:relative;width:{W}px;height:{H}px;overflow:hidden}} .n{{position:absolute;right:4px;bottom:2px;font-size:11px;color:#c00}}</style></head><body>'
       + ''.join(p for _, p in paginas) + '</body></html>')
open(S + f'/{pref}.html', 'w').write(doc)
js = f"""const {{chromium}}=require('playwright');(async()=>{{const b=await chromium.launch();const p=await b.newPage({{viewport:{{width:{int(W)},height:{int(H)}}},deviceScaleFactor:1.5}});
await p.goto('file://{S}/{pref}.html');await p.waitForTimeout(800);const els=await p.$$('.sl');
const nums={[n for n, _ in paginas]};for(let k=0;k<els.length;k++){{await els[k].screenshot({{path:'{S}/{pref}-'+String(nums[k]).padStart(2,'0')+'.png'}});}}await b.close();}})();"""
open(S + '/render_tmp.js', 'w').write(js)
subprocess.run(['node', S + '/render_tmp.js'], check=True, env={**os.environ, 'NODE_PATH': subprocess.check_output(['npm', 'root', '-g']).decode().strip()})
print('ok', [n for n, _ in paginas])
