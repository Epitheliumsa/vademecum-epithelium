# "Ajustes aplicaciones Epithelium": cambios del 1 de octubre de 2026 en el Vademécum y en Visita Comercial (misma plantilla de la capacitación)
from lib_deck import *

VID_V, VID_R, VID_E, VID_C = VID1, VID2, VID3, VID4

def duplicar_portada(src):
    """Copia una portada con video (formas, fondo, animación de reproducción y transición) en una diapositiva nueva."""
    s = prs.slides.add_slide(src.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    mapa = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE): continue
        mapa[rId] = s.part.relate_to(rel._target, rel.reltype, rel.is_external)
    nuevo = copy.deepcopy(src._element)
    for node in nuevo.iter():
        for att in list(node.attrib):
            if att in (qn('r:embed'), qn('r:link'), qn('r:id')) and node.get(att) in mapa: node.set(att, mapa[node.get(att)])
    # se cambia el contenido dentro del mismo spTree (python-pptx guarda referencias a él)
    arbol = s.shapes._spTree
    for ch in list(arbol): arbol.remove(ch)
    for ch in nuevo.find(qn('p:cSld')).find(qn('p:spTree')): arbol.append(ch)
    bg = nuevo.find(qn('p:cSld')).find(qn('p:bg'))
    if bg is not None: s._element.find(qn('p:cSld')).insert(0, bg)
    # resto en el orden del esquema: clrMapOvr, transición, timing, extLst
    for ch in list(s._element):
        if ch.tag != qn('p:cSld'): s._element.remove(ch)
    for ch in nuevo:
        if ch.tag != qn('p:cSld'): s._element.append(ch)
    return s
VID_Z = duplicar_portada(VID_E)

# Video de fondo de cada portada según el tema (video_temas.py)
from pptx.media import Video as _Video
def poner_video(slide, tema):
    mp4, jpg = f'{S}/vid_{tema}.mp4', f'{S}/vid_{tema}.jpg'
    mpart = slide.part.package.get_or_add_media_part(_Video.from_path_or_file_like(mp4, 'video/mp4'))
    ipart, _ = slide.part.get_or_add_image_part(jpg)
    # primero se quitan las relaciones viejas y luego se crean las nuevas (así un rId nuevo no choca con uno viejo)
    viejas = [(rId, rel.reltype) for rId, rel in list(slide.part.rels.items())
              if rel.reltype in (RT.VIDEO, RT.MEDIA) or (rel.reltype == RT.IMAGE and rel.target_part.partname.endswith('.jpg'))]
    for rId, _ in viejas: slide.part.rels.pop(rId)
    cambio = {rId: slide.part.relate_to(ipart if rt == RT.IMAGE else mpart, rt) for rId, rt in viejas}
    for node in slide._element.iter():
        for att in (qn('r:embed'), qn('r:link'), qn('r:id')):
            if node.get(att) in cambio: node.set(att, cambio[node.get(att)])
for _s, _t in [(VID_E, 'etiquetas'), (VID_Z, 'zonas'), (VID_V, 'vademecum'), (VID_R, 'visita'), (VID_C, 'cotizador')]:
    poner_video(_s, _t)
URL_R, URL_V = URL_RUTA, URL_VADE

# ---------------------------------------------------------------- marco de computador
def computador(src, nombre):
    im = Image.open(CAP + src + '.png').convert('RGB')
    B = 36
    W, H = im.width + 2 * B, im.height + 2 * B
    lienzo = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    ImageDraw.Draw(lienzo).rounded_rectangle([0, 0, W - 1, H - 1], radius=40, fill=(28, 28, 30, 255))
    lienzo.paste(im, (B, B))
    lienzo = lienzo.resize((W // 2, H // 2), Image.LANCZOS)
    out = S + '/build/pc_' + nombre + '.png'; lienzo.save(out, optimize=True)
    return out, W, H, B

def poner_pc(slide, src, x, y, ancho, nombre='Computador'):
    path, W, H, B = computador(src, src)
    alto = ancho * H / W
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(ancho), Inches(alto)); pic.name = nombre
    esc = ancho / W
    def mapa(c):
        cx, cy, cw, ch = c
        return (x + (B + cx * 3) * esc, y + (B + cy * 3) * esc, cw * 3 * esc, ch * 3 * esc)
    return pic, mapa, ancho, alto

# ---------------------------------------------------------------- diapositiva genérica: pantallas + tabla referenciada
def diapo(titulo_t, pantallas, filas, encabezado=('#', 'Elemento', 'Qué haces'), anchos=(0.3, 1.25, 3.1), nota=None,
          fila_h=0.5, size=11, regreso=False, y0=1.1, tabla_y=None, tabla_x=None, tabla_w=None, gap=0.3):
    s = nueva(titulo_t); nuevas.append(s); an = Anim(s)
    x = 0.7; pics = []; refs = []
    for pz in pantallas:
        if pz.get('pc'):
            pic, mapa, w, h = poner_pc(s, pz['img'], pz.get('x', x), pz.get('y', y0), pz['ancho'])
        else:
            alto = pz.get('alto', 5.5)
            rec = pz.get('recorte')
            if rec: alto = pz.get('alto', 5.5)
            pic, mapa, w = poner_celular(s, pz['img'], pz.get('x', x), pz.get('y', y0), alto, recorte=rec, nombre='Cel_' + pz['img'])
        pics.append(pic)
        cj = CAJAS.get(pz['img'], {}).get('items', {})
        for r in pz.get('refs', []):
            clave, lado = (r, 'der') if isinstance(r, str) or (isinstance(r, list) and len(r) == 4) else r
            if isinstance(clave, (list, tuple)): caja = clave   # caja manual en px CSS
            else: caja = cj.get(clave)
            if caja and not pz.get('pc'):
                lo, hi = (pz['recorte'] if pz.get('recorte') else (0, CAJAS.get(pz['img'], {}).get('h', 844)))
                y1, y2 = max(caja[1], lo + 4), min(caja[1] + caja[3], hi - 4)
                caja = [caja[0], y1, caja[2], max(12, y2 - y1)] if y2 > lo + 4 else None
            refs.append((caja, lado, mapa))
        x = pz.get('x', x) + w + gap
    # tabla: a la derecha; si queda angosta, se baja por debajo del logo y llega hasta 12.45"
    tx = tabla_x if tabla_x is not None else x + 0.25
    ty = tabla_y if tabla_y is not None else (y0 if XMAX - tx >= 3.9 else 2.65)
    tw = tabla_w if tabla_w is not None else ((XMAX if ty < 2.6 else 12.45) - tx)
    datos = [list(encabezado)] + [[str(i + 1)] + list(f) for i, f in enumerate(filas)]
    gf = tabla(s, tx, ty, tw, list(anchos), datos, alto_fila=fila_h, size=size, centrar=(0,))
    hs = resaltar_filas(s, gf, range(1, len(filas) + 1))
    extra = []
    if nota:
        ny = ty + fila_h * (len(filas) + 1) + 0.15
        extra.append(texto(s, tx, ny, 12.45 - tx, max(0.5, 7.2 - ny), nota, size=11.5, espacio=3))
    an.auto = pics + [gf] + extra
    nref = 0
    for i, h in enumerate(hs):
        formas = []
        if i < len(refs) and refs[i][0]:
            caja, lado, mapa = refs[i]
            formas = llamada(s, i + 1, mapa(caja), lado=lado)
        an.clics.append({'entra': formas + [h], 'sale': [hs[i - 1]] if i else []})
    an.construir()
    if regreso: icono_regreso(s)
    return s

def b(t): return (t, True)


# ---------------------------------------------------------------- categorías con su color (magistrales y materias primas)
def diapo_categorias(titulo_t, items, pie, cols=4, nota=None):
    """items: (nombre, color hex, cantidad). Tarjetas con barra de color, ordenadas como en la app."""
    s = nueva(titulo_t); nuevas.append(s); an = Anim(s)
    filas = (len(items) + cols - 1) // cols
    y0 = 1.1 + (0.75 if nota else 0)
    W, gx, gy = (XMAX - 0.7 - (cols - 1) * 0.12) / cols, 0.12, 0.1
    Hc = min(0.9, (6.75 - y0 - (filas - 1) * gy) / filas)
    tarjetas = []
    if nota:
        tarjetas.append(texto(s, 0.7, 1.05, XMAX - 0.7, 0.65, nota, size=13, color=GRIS, espacio=2))
    for k, (nom, color, n) in enumerate(items):
        x = 0.7 + (k % cols) * (W + gx); y = y0 + (k // cols) * (Hc + gy)
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(W), Inches(Hc)); sh.name = f'Categoria{k + 1}'
        sh.adjustments[0] = 0.12; sh.fill.solid(); sh.fill.fore_color.rgb = rgb('FFFFFF'); sh.line.color.rgb = rgb(color.lstrip('#')); sh.line.width = Pt(2.25); sh.shadow.inherit = False
        tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.42); tf.margin_right = Inches(0.08); tf.margin_top = tf.margin_bottom = Inches(0.03)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT; r = p.add_run(); r.text = nom; r.font.size = Pt(11.5); r.font.bold = True; r.font.color.rgb = rgb('1F2937')
        p = tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT; r = p.add_run(); r.text = f'{n}'; r.font.size = Pt(10); r.font.color.rgb = rgb(GRIS)
        pto = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.12), Inches(y + Hc / 2 - 0.11), Inches(0.22), Inches(0.22)); pto.name = f'Color{k + 1}'
        pto.fill.solid(); pto.fill.fore_color.rgb = rgb(color.lstrip('#')); pto.line.fill.background(); pto.shadow.inherit = False
        tarjetas += [sh, pto]
    pie_tb = fuente(s, pie, y=6.95)
    an.auto = tarjetas + [pie_tb]; an.construir()
    return s

import collections as _col
_DATA = json.load(open('/tmp/vade_main/data.json'))
_CATP = json.load(open('/tmp/vade_main/categorias-productos.json'))
_np = _col.Counter(x['Categoría del Producto'].split(' / ', 1)[-1] for x in _DATA)
CATS_MAGISTRALES = [(k, v, f'{_np[k]} producto' + ('' if _np[k] == 1 else 's')) for k, v in _CATP.items() if _np[k]]
_MP = json.load(open('/tmp/vade_main/materias-primas.json'))
_CATM = json.load(open('/tmp/vade_main/categorias.json'))
_nm = _col.Counter(c for m in _MP for c in {m['Categoria Principal'], m['Categoria Secundaria']} if c)
CATS_MP = [(k, v['color'], f'{_nm[k]} materia' + ('' if _nm[k] == 1 else 's') + ' prima' + ('' if _nm[k] == 1 else 's')) for k, v in _CATM.items() if _nm[k]]


OUT = S + '/202610Ajustes Aplicaciones Epithelium.pptx'
FECHA_AJ = 'Octubre 1 de 2026'
# ================================================================ PORTADA Y ORDEN
ph = PORTADA.shapes.title; tf = ph.text_frame
tf.paragraphs[0].runs[0].text = 'AJUSTES APLICACIONES'
for r in tf.paragraphs[0].runs[1:]: r.text = ''
tf.paragraphs[0].runs[-1].text = 'EPITHELIUM'
for r in tf.paragraphs[0].runs: r.font.size = Pt(44)
sub = [x for x in PORTADA.placeholders if x.placeholder_format.idx == 1][0]
sub.text_frame.paragraphs[0].runs[0].text = 'Vademécum · Visita Comercial'
for r in sub.text_frame.paragraphs[0].runs[1:]: r.text = ''
sub.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
p2 = sub.text_frame.add_paragraph(); r = p2.add_run(); r.text = f'Cambios después de la capacitación · {FECHA_AJ}'; r.font.size = Pt(18)
sub.left, sub.width = Inches(0.9), Inches(7.7)

cont = [x for x in ORDEN.placeholders if x.placeholder_format.idx == 1][0]
cont._element.getparent().remove(cont._element)
puntos = ['Vademécum Epithelium.', 'Visita Comercial.']
PASO_ORDEN, Y_ORDEN = 1.0, 2.2
items_orden = [texto(ORDEN, 1.55, Y_ORDEN + i * PASO_ORDEN, 10.2, 0.6, [[(f'{i + 1}.  ', True), (t, False)]], size=30, color='000000',
                     anchor=MSO_ANCHOR.MIDDLE, espacio=0, nombre=f'Punto{i + 1}') for i, t in enumerate(puntos)]
iconos = sorted([sh for sh in ORDEN.shapes if sh.shape_type == 13], key=lambda x: x.top)
for extra_ic in iconos[len(puntos):]: extra_ic._element.getparent().remove(extra_ic._element)
iconos = iconos[:len(puntos)]
for k, ic in enumerate(iconos):
    ic.left = Inches(0.86); ic.top = Inches(Y_ORDEN + k * PASO_ORDEN + 0.3) - ic.height // 2
for ic, dest in zip(iconos, [VID_V, VID_R]): enlazar(ic, dest)
ao = Anim(ORDEN)
for i, tb in enumerate(items_orden): ao.clics.append({'entra': [tb, iconos[i]]})
ao.construir()
for s_, t, sub_t in [(VID_V, 'Vademécum Epithelium', 'Ajustes · ' + FECHA_AJ), (VID_R, 'Visita Comercial', 'Ajustes · ' + FECHA_AJ)]:
    s_.shapes.title.text_frame.paragraphs[0].runs[0].text = t
    sp = [x for x in s_.placeholders if x.placeholder_format.idx == 1][0]
    sp.text_frame.paragraphs[0].runs[0].text = sub_t; sp.width = Inches(4.4)
    for sh in (s_.shapes.title, sp):
        for r in sh.text_frame.paragraphs[0].runs: r.font.color.rgb = rgb('FFFFFF')

nuevas = []
# ---------------------------------------------------------------- Resumen y claves nuevas
s = nueva('Qué cambió'); nuevas.append(s); an = Anim(s)
ks = [kpi(s, 0.77 + k * 2.69, 1.09, 2.53, 0.98, v, d, f'KPI{k}') for k, (v, d) in enumerate([
    ('3 ajustes', 'En el Vademécum'), ('Reglas nuevas', 'Para programar y cerrar visitas'), ('Claves nuevas', 'En las dos apps'), ('8:00 a. m.', 'Hora límite para programar')])]
tv = texto(s, 0.77, 2.35, 5.6, 0.45, [[('Vademécum Epithelium', True)]], size=16, color=VERDE_OSC)
bv = viñetas(s, 0.77, 2.8, 5.6, 3.9, ['Filtro de Etiqueta para la fuerza de ventas (dato interno).', 'Funciona sin internet después del primer ingreso.',
    'Claves nuevas para usuarios internos y clientes.'], size=14)
tr = texto(s, 6.65, 2.35, 5.8, 0.45, [[('Visita Comercial', True)]], size=16, color=VERDE_OSC)
br = viñetas(s, 6.65, 2.8, 5.8, 4.1, ['Programar: aviso de domingo y festivo, sin días pasados, Trabajo Administrativo con horario.',
    'Orden del día con ▲ Subir / ▼ Bajar hasta las 8:00 a. m.', 'Corregir reportes hasta las 11:59 a. m. del siguiente día hábil.',
    'Acompañamiento de jefes y compañeros.', 'Eliminar visitas futuras (queda la huella).', 'Cumpleaños en los calendarios.',
    'Leads, Maestra y firmas de la solicitud de creación.'], size=13)
an.auto = ks + [tv, bv, tr, br]; an.construir()

s = nueva('Claves nuevas en las dos apps'); nuevas.append(s); an = Anim(s)
pv, mv, wv = poner_celular(s, 'v_login', 0.77, 1.1, 5.2, nombre='Cel_v_login')
pr_, mr, wr = poner_celular(s, 'r_login', 0.77 + wv + 0.3, 1.1, 5.2, nombre='Cel_r_login')
xb = 0.77 + wv + wr + 0.9
tb = viñetas(s, xb, 2.65, 12.45 - xb, 4.0, [
    [('Formato: ', True), ('iniciales del nombre y apellido + dos números.', False)],
    [('Para quién: ', True), ('usuarios internos (Vademécum y Visita Comercial) y clientes del Vademécum.', False)],
    [('Las claves anteriores ya no funcionan: ', True), ('cada uno recibe la suya de forma personal.', False)],
    [('El usuario no cambia: ', True), ('se sigue entrando con el mismo usuario (ej. inicial.apellido).', False)],
    [('Médicos: ', True), ('siguen entrando con el usuario en blanco y la clave de médicos.', False)]], size=15)
an.auto = [pv, pr_]; clics_viñetas(an, tb, 5); an.construir()

# ================================================================ PUNTO 1 · VADEMÉCUM
PUNTO1_INI = len(nuevas)
diapo('Vademécum: filtro de Etiqueta', [
        {'img': 'aj_v_etiqueta', 'alto': 5.2, 'refs': ['etiqueta', 'card']},
        {'img': 'aj_v_detalle', 'alto': 5.2, 'refs': ['etq']},
        {'img': 'aj_v_medico', 'alto': 5.2, 'refs': ['ref']}],
    [('Etiqueta', 'Nuevo filtro: Nuevo, Foco, Transición-Impulso, Portafolio, Cliente…'), ('Resultados', 'Solo los productos de esa etiqueta; se combina con los demás filtros.'),
     ('Ficha', 'La sección "Categorías" ahora se llama "Etiqueta".'), ('Médico o cliente', 'No ve la etiqueta: la ficha termina en Referencia Interna.')],
    anchos=(0.3, 1.1, 2.3), fila_h=0.62, size=10.5,
    nota=[[b('Solo equipo y comerciales: '), 'la etiqueta es información interna. También sale en el portafolio del cliente.']])
diapo('Vademécum sin internet', [
        {'img': 'aj_v_offline', 'alto': 5.2, 'refs': ['card']},
        {'img': 'aj_v_offline_mp', 'alto': 5.2, 'refs': ['card']}],
    [('Productos', 'Se consultan sin señal, con su ficha completa.'), ('Materias primas', 'También quedan guardadas en el celular.'),
     ('Requisito', 'Entrar una vez con internet: ahí se guarda todo.'), ('Con señal', 'Siempre trae lo más reciente.'),
     ('Portafolios de cliente', 'Siguen pidiendo internet (privacidad).')],
    fila_h=0.62, size=11.5, regreso=True,
    nota=[[b('Capturas tomadas en modo avión: '), 'la app abre y muestra productos y materias primas sin conexión.']])

# ================================================================ PUNTO 2 · VISITA COMERCIAL
PUNTO2_INI = len(nuevas)
diapo('Programar: reglas nuevas', [
        {'img': 'aj_domingo', 'alto': 5.2, 'refs': ['dialogo']},
        {'img': 'aj_admin', 'alto': 5.2, 'recorte': (0, 844), 'refs': ['que', 'horas']},
        {'img': 'aj_clasif', 'alto': 5.2, 'recorte': (0, 844), 'refs': ['clasif']}],
    [('Domingo o festivo', 'Sale un aviso y pide confirmar.'), ('Trabajo Administrativo', 'Nuevo nombre del trabajo interno.'),
     ('Horario', '"Todo el día" o con hora de inicio y fin obligatorias.'), ('Clasificación', 'Sale al escoger el cliente.')],
    anchos=(0.3, 1.2, 2.2), fila_h=0.62, size=10.5,
    nota=[[b('Además: '), 'no se programa en días pasados; lo programado hoy después de las 8:00 a. m. queda NO programado; si ya está en el Visiplan, lo confirma; un trabajo administrativo no se repite el mismo día.']])
diapo('Orden del día: Subir y Bajar', [
        {'img': 'aj_orden', 'alto': 3.6, 'x': 0.7, 'y': 1.15, 'recorte': (660, 1000), 'refs': ['insignia', ('mover', 'izq'), ('subir', 'arriba'), 'bajar']},
        {'img': 'aj_orden', 'alto': 2.15, 'x': 0.7, 'y': 4.95, 'recorte': (1500, 1700), 'refs': ['admin']}],
    [('Número', 'Orden programado de la visita.'), ('Orden de visita', 'Posición dentro del día (ej. 1 de 4).'),
     ('▲ Subir', 'La pasa antes.'), ('▼ Bajar', 'La pasa después.'), ('Trabajo Administrativo', 'También lleva número y su horario.')],
    anchos=(0.3, 1.4, 3.0), fila_h=0.6, size=11.5, tabla_x=5.7,
    nota=[[b('Hasta las 8:00 a. m.: '), 'después el orden queda fijo. El orden real se da a medida que se cierran las visitas.']])
diapo('Cierre y corrección de reportes', [
        {'img': 'aj_novisitado', 'alto': 5.2, 'recorte': (0, 844), 'refs': ['obs']},
        {'img': 'aj_corregir', 'alto': 5.2, 'recorte': (300, 1100), 'refs': ['corregir']},
        {'img': 'aj_bloqueada', 'alto': 5.2, 'recorte': (300, 1100), 'refs': ['nota', 'pedir']}],
    [('Observaciones', 'Obligatorias en No visitado (hasta 100 caracteres).'), ('Corregir', 'Editar o pasar de Visitado a No visitado (o al revés).'),
     ('Bloqueada', 'Pasado el plazo queda con candado.'), ('Solicitar corrección', 'Se pide al Gerente General.')],
    anchos=(0.3, 1.2, 2.2), fila_h=0.62, size=10.5,
    nota=[[b('Plazo: '), 'el mismo día y hasta las 11:59 a. m. del siguiente día hábil. Se conserva la hora del primer reporte.']])
diapo('Corrección fuera de plazo', [
        {'img': 'aj_corr_form', 'alto': 5.2, 'refs': ['ayuda', 'motivo', 'enviar']},
        {'img': 'aj_corr_aprobar', 'alto': 5.2, 'recorte': (0, 844), 'refs': ['tarjeta', 'autorizar', 'rechazar']}],
    [('Plazo cerrado', 'La app dice hasta cuándo se podía corregir.'), ('¿Qué hay que corregir?', 'Obligatorio, hasta 100 caracteres.'),
     ('Enviar solicitud', 'Le llega al Gerente General.'), ('Solicitud', 'En "Solicitudes de eliminación y corrección".'),
     ('Autorizar', 'El vendedor tiene 24 horas para corregirla una vez.'), ('Rechazar', 'Con causa obligatoria (50 caracteres).')],
    anchos=(0.3, 1.35, 2.6), fila_h=0.55, size=10.5)
diapo('Acompañamiento', [
        {'img': 'aj_acomp_boton', 'alto': 5.2, 'recorte': (560, 1404), 'refs': ['boton']},
        {'img': 'aj_acomp_form', 'alto': 5.2, 'refs': ['jefes', 'comp', 'enviar']}],
    [('Solicitar acompañamiento', 'Desde la tarjeta de la visita.'), ('Jefes', 'Escoge a quién le pides.'), ('Compañeros', 'También puedes pedirlo a otro vendedor.'),
     ('Enviar', 'Quien la acepta la ve en su programación.')],
    anchos=(0.3, 1.45, 2.9), fila_h=0.6, size=11.5,
    nota=[[b('Jefes: '), 'tienen "🤝 Acompañar" en cualquier visita (no en días pasados). Cada uno reporta la suya y en el historial sale "Visita acompañada" con el reporte del otro.']])
diapo('Eliminar una visita futura', [
        {'img': 'aj_eliminar_dlg', 'alto': 5.2, 'refs': ['dialogo']},
        {'img': 'aj_huella', 'alto': 5.2, 'refs': ['huella']}],
    [('Eliminar', 'Quien la programó la elimina directo (solo días futuros).'), ('Huella', 'Queda una tarjeta pequeña en rojo al final del día; no cuenta en nada.')],
    anchos=(0.3, 1.2, 3.2), fila_h=0.75, size=12,
    nota=[[b('Visitas de hoy o pasadas: '), 'siguen con "Solicitar eliminación"; al rechazarla se pide la razón (50 caracteres) y le sale al vendedor.']])
diapo('Cumpleaños en los calendarios', [
        {'img': 'aj_cumple_agenda', 'alto': 5.2, 'refs': ['semana']},
        {'img': 'aj_cumple_cal', 'alto': 5.2 * 420 / 844, 'recorte': (0, 420), 'refs': ['modal']}],
    [('Semana', 'El día del cumpleaños sale en morado con "Cumple".'), ('Calendario del mes', 'También en el Visiplan; rayado si cae en festivo, sábado o domingo.')],
    anchos=(0.3, 1.3, 3.1), fila_h=0.75, size=12,
    nota=[[b('Ese día: '), 'la app saluda al vendedor al entrar. El cumpleaños ya no es una opción de Novedades.']])
diapo('Leads y Maestra Clientes', [
        {'img': 'aj_leads', 'alto': 5.2, 'refs': ['crear']},
        {'img': 'aj_maestra', 'alto': 5.2, 'refs': ['periodo']},
        {'img': 'aj_maestra_periodo', 'alto': 5.2, 'refs': ['modal']}],
    [('+ Crear Lead', 'Más visible, arriba y en naranja.'), ('Periodo', 'Ícono de calendario en la Maestra.'),
     ('Visitados / No visitados', 'Por día, mes, trimestre, semestre o año.')],
    anchos=(0.3, 1.3, 2.4), fila_h=0.7, size=11,
    nota=[[b('Lead perdida: '), 'la causa es obligatoria (50 caracteres).']])
diapo('Solicitud de creación: firmas en orden', [
        {'img': 'vc_sol_firmada', 'alto': 5.2 * 400 / 844, 'y': 2.0, 'recorte': (300, 700), 'refs': ['firmas']},
        {'img': 'vc_sol_pdf', 'alto': 5.2, 'recorte': (180, 1000), 'refs': ['pdf', 'firmas']}],
    [('Primera firma', 'Coordinador Comercial; Gerencia espera.'), ('PDF aprobado', 'Con todas las firmas, el vendedor lo descarga.'),
     ('Sello', '"✔ FIRMADO ELECTRÓNICAMENTE" en la casilla de cada firma.')],
    anchos=(0.3, 1.2, 2.9), fila_h=0.7, size=11.5,
    nota=[[b('Orden: '), 'primero Coordinador Comercial y después Gerente General. Al rechazar, la causa es obligatoria (50 caracteres).']])
# ---- horarios y reglas clave
s = nueva('Horarios y reglas clave'); nuevas.append(s); an = Anim(s)
gf = tabla(s, 0.77, 1.15, 10.4, [2.4, 8.0], [['Cuándo', 'Qué pasa'],
    ['8:00 a. m.', 'Límite para programar el día y cambiar el orden. Después, lo nuevo queda NO programado.'],
    ['Día hábil anterior', 'Desde ahí se confirma lo del Visiplan (después de las 8:00 a. m. del día queda NO programado).'],
    ['2.º día hábil, 11:59 p. m.', 'Último momento para editar el Visiplan del mes.'],
    ['11:59 a. m. siguiente día hábil', 'Último momento para corregir un reporte; después, solicitud al Gerente General.'],
    ['24 horas', 'Tiempo para corregir una visita con corrección autorizada.'],
    ['Cada 30 segundos', 'La app se sincroniza sola (solicitudes, acompañamientos, aprobaciones).']], alto_fila=0.55, size=12.5)
hs = resaltar_filas(s, gf, range(1, 7))
tb = viñetas(s, 0.77, 5.3, 11.6, 1.6, [
    [('Obligatorios: ', True), ('causa al rechazar (50), causa de Lead perdida (50), detalle de Cita médica (50), observaciones de No visitado (100). Vacaciones: detalle opcional.', False)],
    [('Más ágil: ', True), ('ícono más grande; recuadros que solo se cierran con X, Cancelar o atrás; el botón atrás del Android vuelve al inicio; nombres en Nombre Propio.', False)]], size=13)
an.auto = [gf]
for i, h in enumerate(hs): an.clics.append({'entra': [h], 'sale': [hs[i - 1]] if i else []})
clics_viñetas(an, tb, 2); an.construir()
icono_regreso(s)

# ================================================================ CIERRE
cierre = prs.slides.add_slide([l for l in prs.slide_layouts if l.part.partname.endswith('slideLayout3.xml')][0])
for ph_ in list(cierre.placeholders): ph_._element.getparent().remove(ph_._element)

orden = [PORTADA, ORDEN] + nuevas[:PUNTO1_INI] + [VID_V] + nuevas[PUNTO1_INI:PUNTO2_INI] + [VID_R] + nuevas[PUNTO2_INI:] + [cierre]
lst = prs.slides._sldIdLst
ids = {prs.part.related_part(e.rId): e for e in list(lst)}
conservar = {s_.part for s_ in orden}
for e in list(lst):
    if prs.part.related_part(e.rId) not in conservar: prs.part.drop_rel(e.rId)
    lst.remove(e)
for s_ in orden: lst.append(ids[s_.part])
for sl in orden:
    xml = etree.tostring(sl._element).decode()
    for rId, rel in list(sl.part.rels.items()):
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE): continue
        if f'"{rId}"' not in xml: sl.part.rels.pop(rId)
for ext in prs.part._element.findall('.//' + qn('p:ext')):
    if b'sectionLst' in etree.tostring(ext): ext.getparent().remove(ext)
vp = prs.part.part_related_by(RT.VIEW_PROPS)
vx = etree.fromstring(vp.blob)
for sl in vx.iter('{%s}sldLst' % NS['p']): sl.getparent().remove(sl)
vp._blob = etree.tostring(vx, xml_declaration=True, encoding='UTF-8', standalone=True)
for rId, rel in list(vp.rels.items()):
    if rel.reltype == RT.SLIDE: vp.rels.pop(rId)
for sl in nuevas + [ORDEN, cierre]: poner_transicion(sl)
for sl in REGRESOS: bloquear_avance(sl)
prs.save(OUT)
print('ok', len(orden), 'diapositivas')
