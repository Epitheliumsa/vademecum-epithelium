#!/usr/bin/env python3
"""Sello de seguridad y aviso de derechos para lo que se entrega a usuarios EXTERNOS (PDF y Excel).

Regla del proyecto: todo PDF o Excel que se envíe a un cliente, médico o invitado pasa por esta herramienta.
Los usuarios internos (equipo y comerciales) no llevan sello.

Uso:
    python3 herramientas/marca_externos.py ARCHIVO "Nombre del cliente" [--salida RUTA] [--sin-pie]

PDF  : sello redondo muy tenue (mismo de la app) repetido en todas las páginas y, si no se usa --sin-pie, el aviso de
       derechos reservados al pie. Usa --sin-pie cuando el documento ya trae el aviso escrito.
Excel: hoja "Aviso legal" al inicio, franja con el nombre del cliente sobre la tabla de cada hoja, encabezado y pie
       de impresión con el aviso y las propiedades del archivo. Sin líneas de cuadrícula.
Requisitos: PyMuPDF (pip install pymupdf), openpyxl y playwright con Chromium.
"""
import argparse
import datetime
import os
import re
import sys
import tempfile
import urllib.parse

COLOR = '0,107,79'   # verde institucional


def texto_legal(anio=None):
    anio = anio or datetime.date.today().year
    return (f'© {anio} Epithelium S.A. Todos los derechos reservados.',
            'La presente aplicación, sus contenidos, bases de datos, código fuente y material gráfico son propiedad '
            'exclusiva de Epithelium S.A. Queda estrictamente prohibida su reproducción, adaptación, distribución, '
            'comercialización, transmisión o divulgación, total o parcial, por cualquier medio, sin la autorización '
            'previa, expresa y por escrito de Epithelium S.A. Información confidencial y de uso exclusivo para '
            'usuarios autorizados.')


# ---------- sello ----------
def lineas_sello(nombre):
    """Reparte el nombre en líneas (igual que la app): letra más pequeña si es largo, recorta con "…"."""
    def partir(maximo):
        out, actual = [], ''
        for p in nombre.split():
            if actual and len(actual + ' ' + p) > maximo:
                out.append(actual)
                actual = p
            else:
                actual = f'{actual} {p}'.strip()
        if actual:
            out.append(actual)
        return out
    lineas, tam, paso = partir(18), 9.5, 13
    if len(lineas) > 4:
        lineas, tam, paso = partir(25), 7.6, 10.5
    if len(lineas) > 5:
        lineas = lineas[:5]
        lineas[4] = lineas[4][:-3] + '…'
    return lineas, tam, paso


def svg_sello(nombre, alfa=0.07):
    esc = lambda t: t.replace('&', '&amp;').replace('<', '&lt;')
    lineas, tam, paso = lineas_sello(nombre)
    y0 = 171 - (len(lineas) - 1) * paso / 2
    texto = ''.join(f'<text x="170" y="{y0 + k * paso}" font-size="{tam}" font-weight="bold">{esc(l)}</text>'
                    for k, l in enumerate(lineas))
    c = f'rgba({COLOR},{alfa})'
    return (f'<svg xmlns="http://www.w3.org/2000/svg" xml:space="preserve" width="340" height="340" viewBox="0 0 340 340">'
            f'<g transform="rotate(-18 170 170)" fill="{c}" stroke="{c}" font-family="Arial,Helvetica,sans-serif" text-anchor="middle">'
            '<defs><path id="c" d="M170,170 m-76,0 a76,76 0 1,1 152,0 a76,76 0 1,1 -152,0"/></defs>'
            '<circle cx="170" cy="170" r="108" fill="none" stroke-width="3.2"/><circle cx="170" cy="170" r="102" fill="none" stroke-width="1"/>'
            '<circle cx="170" cy="170" r="56" fill="none" stroke-width="1.6"/><g stroke="none">'
            '<text font-size="12" font-weight="bold" letter-spacing="1.2" word-spacing="3"><textPath href="#c" startOffset="50%">'
            'EPITHELIUM S.A. • USO EXCLUSIVO • CONFIDENCIAL •</textPath></text>'
            f'{texto}<text x="170" y="{y0 + len(lineas) * paso + 4}" font-size="7" letter-spacing="0.4">© Epithelium S.A.</text>'
            '</g></g></svg>')


# ---------- PDF ----------
def _pdf_capa(ancho_pt, alto_pt, nombre, pie, alfa):
    """Genera (con Chromium) una página transparente con el sello en mosaico y, opcional, el aviso al pie."""
    from playwright.sync_api import sync_playwright
    titulo, cuerpo = texto_legal()
    lado = round(ancho_pt * 0.42)
    fondo = 'data:image/svg+xml,' + urllib.parse.quote(svg_sello(nombre, alfa))
    aviso = (f'<p class="aviso"><b>{titulo}</b> {cuerpo}</p>' if pie else '')
    html = f"""<!DOCTYPE html><html lang="es"><meta charset="utf-8"><style>
      @page {{ size: {ancho_pt}pt {alto_pt}pt; margin: 0 }}
      html, body {{ margin: 0; background: transparent }}
      .sello {{ position: fixed; inset: 0; background: url("{fondo}") repeat; background-size: {lado}pt {lado}pt }}
      .aviso {{ position: fixed; left: 28pt; right: 28pt; bottom: 12pt; margin: 0; font: 6.4pt/1.3 Arial, sans-serif;
                color: #5a6f68; text-align: justify; hyphens: auto }}
    </style><div class="sello"></div>{aviso}"""
    with sync_playwright() as p:
        try:
            nav = p.chromium.launch()
        except Exception:
            nav = p.chromium.launch(executable_path=os.environ.get('CHROMIUM_PATH', '/opt/pw-browsers/chromium'))
        pag = nav.new_page()
        pag.set_content(html)
        ruta = tempfile.mktemp(suffix='.pdf')
        pag.pdf(path=ruta, width=f'{ancho_pt / 72}in', height=f'{alto_pt / 72}in', print_background=True,
                margin={'top': '0', 'right': '0', 'bottom': '0', 'left': '0'})
        nav.close()
    return ruta


def aplicar_a_pdf(entrada, salida, nombre, pie=True, alfa=0.07):
    import pymupdf
    doc = pymupdf.open(entrada)
    capas = {}
    for pagina in doc:
        ancho, alto = round(pagina.rect.width), round(pagina.rect.height)
        if (ancho, alto) not in capas:
            capas[(ancho, alto)] = pymupdf.open(_pdf_capa(ancho, alto, nombre, pie, alfa))
        pagina.show_pdf_page(pagina.rect, capas[(ancho, alto)], 0, overlay=True)
    doc.set_metadata({**doc.metadata, 'subject': f'Uso exclusivo de {nombre}', 'keywords': 'Epithelium S.A. · Confidencial'})
    doc.save(salida, garbage=3, deflate=True)
    doc.close()
    return salida


# ---------- Excel ----------
def aplicar_a_excel(entrada, salida, nombre):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter, range_boundaries

    titulo, cuerpo = texto_legal()
    corto = f'© {datetime.date.today().year} Epithelium S.A. · Uso exclusivo de {nombre} · Confidencial'
    prohibicion = 'Prohibida su reproducción, adaptación o divulgación sin autorización previa y por escrito de Epithelium S.A.'
    verde = PatternFill('solid', fgColor='E3F2EC')

    wb = load_workbook(entrada)
    for ws in list(wb.worksheets):
        # franja con el nombre del cliente sobre la tabla (solo si la hoja es una tabla simple que empieza en A1)
        tablas = list(ws.tables.values())
        simple = (len(tablas) == 1 and tablas[0].ref.startswith('A1') and not ws.merged_cells.ranges
                  and not ws.conditional_formatting and not ws.data_validations.dataValidation)
        if simple:
            t = tablas[0]
            c1, f1, c2, f2 = range_boundaries(t.ref)
            ws.move_range(f'A1:{get_column_letter(max(c2, ws.max_column))}{max(f2, ws.max_row)}', rows=2)
            nuevo = f'{get_column_letter(c1)}3:{get_column_letter(c2)}{f2 + 2}'
            t.ref = nuevo
            if t.autoFilter is not None:
                t.autoFilter.ref = nuevo
            if ws.freeze_panes:
                ws.freeze_panes = re.sub(r'(\d+)$', lambda m: str(int(m.group(1)) + 2), ws.freeze_panes)
            ws['A1'] = corto
            ws['A1'].font = Font(bold=True, size=11, color='006B4F')
            ws['A2'] = prohibicion
            ws['A2'].font = Font(italic=True, size=9, color='5A6F68')
            for col in range(1, c2 + 1):
                ws.cell(row=1, column=col).fill = verde
                ws.cell(row=2, column=col).fill = verde
        ws.sheet_view.showGridLines = False
        ws.oddHeader.center.text = corto
        ws.oddHeader.center.size = 8
        ws.oddFooter.center.text = prohibicion
        ws.oddFooter.center.size = 7

    aviso = wb.create_sheet('Aviso legal', 0)
    aviso.sheet_view.showGridLines = False
    aviso.column_dimensions['A'].width = 110
    filas = [(titulo, Font(bold=True, size=14, color='006B4F')),
             (f'Documento de uso exclusivo de: {nombre}', Font(bold=True, size=12)),
             (f'Entregado el {datetime.date.today():%d/%m/%Y}', Font(size=10, color='5A6F68')),
             ('', None), (cuerpo, Font(size=11))]
    for i, (txt, fuente) in enumerate(filas, 1):
        c = aviso.cell(row=i, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='justify')
        if fuente:
            c.font = fuente
    aviso.row_dimensions[5].height = 110
    aviso.oddHeader.center.text = corto
    aviso.oddFooter.center.text = prohibicion

    wb.properties.subject = f'Uso exclusivo de {nombre}'
    wb.properties.keywords = 'Epithelium S.A.; Confidencial'
    wb.properties.comments = f'{titulo} {cuerpo}'
    wb.active = 0
    wb.save(salida)
    return salida


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('archivo')
    ap.add_argument('nombre', help='nombre del cliente o usuario externo')
    ap.add_argument('--salida')
    ap.add_argument('--sin-pie', action='store_true', help='PDF: no agrega el aviso al pie (el documento ya lo trae)')
    a = ap.parse_args()
    base, ext = os.path.splitext(a.archivo)
    salida = a.salida or f'{base} - con sello{ext}'
    if ext.lower() == '.pdf':
        aplicar_a_pdf(a.archivo, salida, a.nombre, pie=not a.sin_pie)
    elif ext.lower() in ('.xlsx', '.xlsm'):
        aplicar_a_excel(a.archivo, salida, a.nombre)
    else:
        sys.exit('Solo se admiten archivos PDF y Excel (.xlsx)')
    print(salida)


if __name__ == '__main__':
    main()
